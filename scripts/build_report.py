#!/usr/bin/env python3
"""Turn the scraped CSV (../data/) into a formatted Excel workbook + print a summary."""
import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

def source_stem():
    cands = [p for p in DATA.glob("*.xlsx") if "_with_final_bids" not in p.name]
    if not cands:
        raise SystemExit(f"No source .xlsx found in {DATA}")
    return max(cands, key=lambda p: p.stat().st_mtime).stem

STEM = source_stem()
CSV  = DATA / f"{STEM}_with_final_bids.csv"
XLSX = DATA / f"{STEM}_with_final_bids.xlsx"

MONEY = {"Minimum Bid","Debt Amount","Final Bid","Page Min Bid (check)",
         "Delta vs Min Bid","Delta vs Debt"}
INTS  = {"Num Bids"}
PCTS  = {"% Over Min Bid"}

def num(v):
    if v is None or v == "": return None
    try: return float(v) if "." in str(v) else int(v)
    except ValueError: return v

def load():
    with open(CSV, encoding="utf-8-sig") as f:
        r = csv.DictReader(f)
        return r.fieldnames, list(r)

def build_xlsx(headers, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Auctions"
    hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(bold=True, color="FFFFFF")
    postponed_fill = PatternFill("solid", fgColor="FCE4D6")
    border = Border(bottom=Side(style="thin", color="D9D9D9"))
    for c, name in enumerate(headers, 1):
        cell = ws.cell(1, c, name); cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for ri, row in enumerate(rows, 2):
        status = (row.get("Page Status") or "").lower()
        shade = any(k in status for k in ("postpone","cancel","stay","withdraw"))
        for c, name in enumerate(headers, 1):
            val = num(row[name]) if name in MONEY | INTS | PCTS else row[name]
            cell = ws.cell(ri, c, val)
            if name in MONEY: cell.number_format = '$#,##0'
            elif name in PCTS and isinstance(val, (int, float)): cell.number_format = '0.0"%"'
            cell.border = border
            if shade: cell.fill = postponed_fill
    widths = {"Auction ID":10,"Status":9,"Minimum Bid":12,"Open Date/Time":16,
              "Close Date/Time":16,"Attorney":24,"Debt Amount":13,"Book/Writ":10,
              "OPA":11,"Address":34,"County":11,"Sale Type":18,"Final Bid":12,
              "Page Status":17,"Num Bids":9,"Auction Closed":19,
              "Page Min Bid (check)":14,"Delta vs Min Bid":15,"% Over Min Bid":13,
              "Delta vs Debt":14,"Fetch Note":22}
    for c, name in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    return wb

def summarize(rows):
    def isint(v):
        try: return int(v)
        except (TypeError, ValueError): return None
    sold  = [r for r in rows if (r.get("Page Status") or "").lower().startswith("sold")]
    plaintiff = [r for r in sold if "plaintiff" in r["Page Status"].lower()]
    third = [r for r in sold if r["Page Status"].strip().lower() == "sold"]
    post  = [r for r in rows if "postpone" in (r.get("Page Status") or "").lower()]
    stayed= [r for r in rows if "stay" in (r.get("Page Status") or "").lower()]
    other = [r for r in rows if r not in sold and r not in post and r not in stayed]
    fb = lambda r: isint(r["Final Bid"]); mb = lambda r: isint(r["Minimum Bid"])
    sold_fb = [r for r in sold if fb(r) is not None]
    tot_final = sum(fb(r) for r in sold_fb)
    tot_min   = sum(mb(r) for r in sold_fb if mb(r) is not None)
    print("\n" + "="*60)
    print("SUMMARY —", STEM)
    print("="*60)
    print(f"Total auctions on list ....... {len(rows)}")
    print(f"  Sold ....................... {len(sold)}")
    print(f"    - to Plaintiff (bank) .... {len(plaintiff)}")
    print(f"    - to Third Party (buyer) . {len(third)}")
    print(f"  Postponed .................. {len(post)}")
    print(f"  Stayed ..................... {len(stayed)}")
    print(f"  Other/blank ................ {len(other)}")
    if sold_fb:
        print(f"\nSold total final bids ........ ${tot_final:,}")
        print(f"Sold total minimum bids ...... ${tot_min:,}")
        print(f"Total premium over minimum ... ${tot_final-tot_min:,}")
        prem = [((fb(r)-mb(r))/mb(r)*100) for r in sold_fb if mb(r)]
        print(f"Avg % over minimum ........... {sum(prem)/len(prem):.1f}%")
    notes = [r for r in rows if r.get("Fetch Note")]
    if notes:
        print(f"\nRows needing attention ({len(notes)}):")
        for r in notes: print(f"  {r['Auction ID']}: {r['Fetch Note']}")

if __name__ == "__main__":
    headers, rows = load()
    build_xlsx(headers, rows).save(XLSX)
    print("XLSX written:", XLSX, f"({len(rows)} rows)")
    summarize(rows)
